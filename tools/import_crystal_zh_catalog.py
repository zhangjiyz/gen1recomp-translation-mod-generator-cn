#!/usr/bin/env python3
"""Convert pokecrystal_cn_build's text.xlsx into a Crystal zh-Hans catalog.

The workbook never enters a release archive. A user's own Crystal ROM
extraction joins its current English text to these source-labelled rows and
thereby resolves translations to that ROM's bank:address keys during the
private build. Ambiguous prose remains English unless every candidate ships
the same translation.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re


SCHEMA = "gen1recomp-translation-mods/crystal-zh-labels"
SHEETS = tuple(f"文{index}" for index in range(1, 10))
CONTROL_MARKER = re.compile(r"【(\d+)】")
TOKENS = {
    "<PLAYER>": "{PLAYER}",
    "<PLAY_G>": "{PLAYER}",
    "<RIVAL>": "{RIVAL}",
    "<USER>": "{USER}",
    "<TARGET>": "{TARGET}",
    "<ENEMY>": "{ENEMY}",
    "<……>": "……",
    "<SCROLL>": "",
}


def _control_value(control: str) -> str:
    if control.startswith("text_ram "):
        return "{STRBUF}"
    if control.startswith("text_decimal "):
        return "{NUM}"
    # TX_LOW/SCROLL/PAUSE/WAIT/DAY and sound commands print no glyphs in
    # RomExtractorGen2.decodeGen2Text. LINE_CR only selects NEXT instead of
    # LINE/CONT, which is handled while joining lines below.
    return ""


def _replace_controls(value: str, controls: list[str]) -> str:
    def replace(match: re.Match[str]) -> str:
        index = int(match.group(1))
        if index >= len(controls):
            raise ValueError(f"control marker {match.group(0)!r} has no matching control")
        return _control_value(controls[index])

    value = CONTROL_MARKER.sub(replace, value).replace("|", "")
    for source, target in TOKENS.items():
        value = value.replace(source, target)
    unknown = re.findall(r"<[^>]+>", value)
    if unknown:
        raise ValueError(f"unknown Crystal text token(s): {unknown!r}")
    return value


def _runtime_text(lines: list[str], controls: list[str], next_cr: bool) -> str:
    output: list[str] = []
    line_count = 0
    pending_paragraph = False
    for raw in lines:
        if raw == "":
            if output:
                line_count = 0
                pending_paragraph = True
            continue
        line = _replace_controls(raw, controls)
        if not output:
            output.append(line)
        elif pending_paragraph:
            output.extend(("\f", line))
        elif line_count == 1 or next_cr:
            output.extend(("\n", line))
        else:
            output.extend(("\v", line))
        line_count += 1
        pending_paragraph = False
    return "".join(output)


def _mapping_rows(workbook) -> list[dict[str, str]]:
    sheet = workbook["标"]
    mappings: list[dict[str, str]] = []
    # Match the upstream importer's range(2, ws.max_row): its final row is a
    # workbook sentinel, not a mapping. Iterating is important here:
    # read-only openpyxl cell() access restarts the XML stream per cell.
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row - 1, values_only=True):
        dmap = str(row[0] or "")
        dlabel = str(row[2] or "")
        olabel = str(row[3] or "")
        if not olabel:
            continue
        mappings.append({"map": dmap, "label": dlabel, "block": olabel})
    return mappings


def _text_blocks(workbook) -> dict[str, tuple[str, str]]:
    result: dict[str, tuple[str, str]] = {}
    for sheet_name in SHEETS:
        sheet = workbook[sheet_name]
        olabel = ""
        english_lines: list[str] = []
        chinese: list[str] = []
        controls: list[str] = []

        def finish() -> None:
            nonlocal english_lines, chinese, controls
            if not olabel:
                return
            while english_lines and english_lines[-1] == "":
                english_lines.pop()
            while chinese and chinese[-1] == "":
                chinese.pop()
            next_cr = bool(controls and controls[0] == "LINE_CR")
            if next_cr:
                controls = controls[1:]
            result[olabel] = (
                _runtime_text(english_lines, controls, next_cr),
                _runtime_text(chinese, controls, next_cr),
            )

        for row in sheet.iter_rows(values_only=True):
            english = str(row[0] or "")
            control = str(row[9] or "")
            if "英文" in english or "结束" in english:
                finish()
                if "英文" in english:
                    olabel = control
                    english_lines = []
                    chinese = []
                    controls = []
                else:
                    olabel = ""
                continue
            if not olabel:
                continue
            english_lines.append(english)
            chinese.append(str(row[4] or ""))
            if control:
                controls.append(control)
        finish()
    return result


def build_catalog(workbook_path: Path, source_commit: str, build_commit: str) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise SystemExit("openpyxl is required to import Crystal text.xlsx") from error
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    mappings = _mapping_rows(workbook)
    blocks = _text_blocks(workbook)
    entries = []
    for mapping in mappings:
        block = blocks.get(mapping["block"])
        if block is None:
            continue
        english, translation = block
        if not english or not translation:
            continue
        qid = f'{mapping["map"]}:{mapping["label"]}:{mapping["block"]}'
        entries.append({
            "qid": qid,
            **mapping,
            "english": english,
            "translation": translation,
        })
    entries.sort(key=lambda row: row["qid"])
    missing_blocks = sorted({row["block"] for row in mappings} - set(blocks))
    if missing_blocks:
        raise ValueError(f"mapped workbook blocks are missing: {missing_blocks[:10]!r}")
    return {
        "schema": SCHEMA,
        "version": 1,
        "language": "zh-Hans",
        "source": {
            "translation_repository": "https://github.com/SnDream/pokecrystal_cn",
            "translation_commit": source_commit,
            "build_repository": "https://github.com/SnDream/pokecrystal_cn_build",
            "build_commit": build_commit,
            "workbook": "text.xlsx",
            "workbook_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        },
        "counts": {
            "mapping_rows": len(mappings),
            "mapped_labels": sum(bool(row["label"]) for row in mappings),
            "workbook_text_blocks": len(blocks),
            "translated_rows": len(entries),
            "empty_or_unmapped_rows": len(mappings) - len(entries),
        },
        "entries": entries,
        "notice": (
            "Fan-translation text has no explicit redistribution license in the pinned "
            "source repositories; obtain author permission before public distribution."
        ),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--source-commit", required=True)
    parser.add_argument("--build-commit", required=True)
    args = parser.parse_args()
    catalog = build_catalog(args.workbook.resolve(), args.source_commit, args.build_commit)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(catalog["counts"], ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
